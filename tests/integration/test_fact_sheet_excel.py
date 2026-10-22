import json
from datetime import date
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook
from reportlab.pdfgen.canvas import Canvas

from bd_packet_builder.contracts import Citation, IdentifierValue
from bd_packet_builder.evidence import extract_evidence
from bd_packet_builder.fact_sheet import Fact, FactSheet, FactSheetError, run_fact_sheet_check
from bd_packet_builder.fact_sheet_excel import export_workbook
from bd_packet_builder.inventory import build_inventory

AS_OF = date(2026, 10, 21)
DRAFT = Path("outputs/draft.xlsx")


@pytest.fixture
def workbook_draft(tmp_path: Path) -> Path:
    source = tmp_path / "bd_agreements/supporting_data/synthetic.pdf"
    source.parent.mkdir(parents=True)
    canvas = Canvas(str(source))
    canvas.drawString(40, 700, "Synthetic CRD: 001234")
    canvas.save()
    inventory = build_inventory(tmp_path)
    page = extract_evidence(tmp_path, inventory).documents[0].pages[0]
    fact = Fact(
        fact_id="crd-1",
        key="firm.crd",
        category="firm_identity",
        value=IdentifierValue(kind="identifier", value="001234"),
        owner="Synthetic Owner",
        reviewed_by="Synthetic Reviewer",
        reviewed_on=date(2026, 10, 1),
        review_due=date(2027, 10, 1),
        effective_from=date(2026, 1, 1),
        citations=(
            Citation(
                source_id=page.source_id, page=1, start=0, end=len(page.text), quote=page.text
            ),
        ),
    )
    sheet = FactSheet(version_id="example-v1", steward="Synthetic Steward", facts=(fact,))
    (tmp_path / DRAFT).parent.mkdir()
    (tmp_path / DRAFT).write_bytes(export_workbook(sheet, inventory))
    return tmp_path


def check(root: Path) -> tuple[Path, dict]:
    report = run_fact_sheet_check(root, DRAFT, as_of=AS_OF)
    return report, json.loads((report / "fact-sheet-check.json").read_text())


def edit(root: Path, sheet: str, cell: str, value: object) -> None:
    workbook = load_workbook(root / DRAFT)
    workbook[sheet][cell] = value
    workbook.save(root / DRAFT)
    workbook.close()


def test_workbook_round_trip_preserves_ids_evidence_and_editability(workbook_draft: Path) -> None:
    root = workbook_draft
    before = (root / DRAFT).read_bytes()
    report, data = check(root)
    assert data["facts_with_exceptions"] == 0
    assert data["items"][0]["fact"]["value"]["value"] == "001234"
    assert data["items"][0]["check"]["resolved_citations"] == 1
    assert data["items"][0]["check"]["usable_for_filling"] is False
    assert (root / DRAFT).read_bytes() == before
    workbook = load_workbook(report / "fact-sheet.xlsx")
    assert workbook.sheetnames == ["Instructions", "Details", "Facts", "Evidence", "Checks"]
    assert workbook["Facts"]["E2"].value == "001234"
    assert workbook["Facts"]["E2"].data_type == "s"
    assert workbook["Facts"].freeze_panes == "F2"
    assert workbook["Facts"].data_validations.count >= 2
    assert workbook["Checks"]["B2"].value == "No mechanical exceptions; NOT APPROVED"
    workbook.save(root / DRAFT)
    workbook.close()
    assert check(root)[1]["facts_with_exceptions"] == 0
    edit(root, "Facts", "J2", "2026-10-20")
    assert check(root)[1]["items"][0]["check"]["issues"] == ["review_overdue"]


@pytest.mark.parametrize("value", [1234, True, "=1+1"])
def test_identifiers_are_never_coerced_or_calculated(workbook_draft: Path, value: object) -> None:
    edit(workbook_draft, "Facts", "E2", value)
    with pytest.raises(FactSheetError, match="Facts.*E2"):
        check(workbook_draft)
    assert not list((workbook_draft / "outputs").glob("fact-sheet-check-*"))


@pytest.mark.parametrize(
    "cell,value",
    [
        ("A2", "bd_agreements/Completed/signed.pdf"),
        ("A2", "../outside.pdf"),
        ("A2", "bd_agreements/supporting_data/missing.pdf"),
        ("C2", 999),
        ("D2", "Made-up quotation"),
        ("E2", "0" * 64),
    ],
)
def test_bad_evidence_is_reported_without_losing_fact(
    workbook_draft: Path,
    cell: str,
    value: object,
) -> None:
    # Evidence: fact ID, source path, page, quotation, source hash.
    if cell == "A2":
        cell = "B2"
    edit(workbook_draft, "Evidence", cell, value)
    _, data = check(workbook_draft)
    assert len(data["items"]) == 1
    assert data["items"][0]["check"]["issues"]
    assert data["items"][0]["check"]["usable_for_filling"] is False
    assert "workbook_evidence" in data


def test_empty_evidence_is_explicit(workbook_draft: Path) -> None:
    wb = load_workbook(workbook_draft / DRAFT)
    wb["Evidence"].delete_rows(2)
    wb.save(workbook_draft / DRAFT)
    wb.close()
    assert check(workbook_draft)[1]["items"][0]["check"]["issues"] == ["missing_evidence"]


@pytest.mark.parametrize(
    "sheet,cell,value",
    [
        ("Facts", "A1", "Renamed column"),
        ("Details", "A5", "approved"),
        ("Evidence", "A2", "unknown-fact"),
        ("Facts", "I2", "10/1/26"),
    ],
)
def test_invalid_structure_is_actionable_without_logging_values(
    workbook_draft: Path,
    sheet: str,
    cell: str,
    value: object,
) -> None:
    edit(workbook_draft, sheet, cell, value)
    with pytest.raises(FactSheetError) as error:
        check(workbook_draft)
    assert str(value) not in str(error.value)
    assert sheet in str(error.value)


def test_formulas_are_not_executed_and_output_text_cannot_become_a_formula(
    workbook_draft: Path,
) -> None:
    edit(workbook_draft, "Facts", "G2", '=HYPERLINK("https://example.invalid")')
    with pytest.raises(FactSheetError, match="formula"):
        check(workbook_draft)
    # Intentional literal text beginning with '=' stays text on export.
    root = workbook_draft
    edit(root, "Facts", "G2", "Safe Owner")
    report, _ = check(root)
    from bd_packet_builder.fact_sheet_excel import import_workbook

    sheet, _, _ = import_workbook((root / DRAFT).read_bytes(), root, build_inventory(root))
    fact = sheet.facts[0].model_copy(update={"owner": "=1+1"})
    updated = sheet.model_copy(update={"facts": (fact,)})
    wb = load_workbook(BytesIO(export_workbook(updated, build_inventory(root))))
    assert wb["Facts"]["G2"].data_type == "s"
    assert wb["Facts"]["G2"].value == "=1+1"
    wb.close()
    assert (report / "fact-sheet.xlsx").is_file()


def test_xlsx_zip_expansion_is_bounded(workbook_draft: Path) -> None:
    with ZipFile(workbook_draft / DRAFT, "a", ZIP_DEFLATED) as archive:
        archive.writestr("oversized.xml", b" " * (21 * 1024 * 1024))
    with pytest.raises(FactSheetError, match="expanded"):
        check(workbook_draft)


def test_corrupt_workbook_is_a_safe_failure(workbook_draft: Path) -> None:
    (workbook_draft / DRAFT).write_bytes(b"not a workbook")
    with pytest.raises(FactSheetError, match="workbook"):
        check(workbook_draft)


def test_cli_prints_excel_deliverable(
    workbook_draft: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    from bd_packet_builder.cli import main

    assert (
        main(
            [
                "fact-sheet-check",
                "--root",
                str(workbook_draft),
                "--source",
                str(DRAFT),
                "--as-of",
                AS_OF.isoformat(),
            ]
        )
        == 0
    )
    assert "fact-sheet.xlsx" in capsys.readouterr().out


def test_business_maintainer_instructions_require_neither_code_nor_separate_auditor(
    workbook_draft: Path,
) -> None:
    edit(workbook_draft, "Facts", "H2", "Synthetic Owner")
    report, data = check(workbook_draft)
    assert data["facts_with_exceptions"] == 0
    workbook = load_workbook(report / "fact-sheet.xlsx")
    instructions = "\n".join(
        str(row[0]) for row in workbook["Instructions"].iter_rows(values_only=True)
    )
    workbook.close()
    assert "No separate auditor" in instructions
    assert "No code or terminal" in instructions
    assert "business owner is responsible" in instructions
    assert "plausible typo" in instructions
    assert "run fact-sheet-check" not in instructions
    assert "clear its hash" not in instructions


def test_excel_date_cells_and_new_evidence_without_hash_are_resolved(workbook_draft: Path) -> None:
    from datetime import datetime

    edit(workbook_draft, "Facts", "I2", datetime(2026, 10, 1))
    edit(workbook_draft, "Evidence", "E2", None)
    report, data = check(workbook_draft)
    assert data["facts_with_exceptions"] == 0
    workbook = load_workbook(report / "fact-sheet.xlsx")
    assert len(workbook["Evidence"]["E2"].value) == 64
    workbook.close()


def test_editing_a_value_does_not_make_it_semantically_verified(workbook_draft: Path) -> None:
    edit(workbook_draft, "Facts", "E2", "009999")
    _, data = check(workbook_draft)
    assert data["items"][0]["fact"]["value"]["value"] == "009999"
    assert "semantic_support_unverified" in data["flags"]
    assert data["items"][0]["check"]["usable_for_filling"] is False


def test_ambiguous_quotation_requires_more_context(workbook_draft: Path) -> None:
    edit(workbook_draft, "Evidence", "D2", "0")
    assert (
        "citation_quote_missing_or_ambiguous"
        in check(workbook_draft)[1]["items"][0]["check"]["issues"]
    )


def test_unknown_evidence_is_preserved_in_excel_output(workbook_draft: Path) -> None:
    edit(workbook_draft, "Evidence", "B2", "bd_agreements/supporting_data/missing.pdf")
    report, _ = check(workbook_draft)
    wb = load_workbook(report / "fact-sheet.xlsx")
    assert wb["Evidence"]["B2"].value == "bd_agreements/supporting_data/missing.pdf"
    assert "citation_unresolved" in wb["Checks"]["B2"].value
    wb.close()


def test_failed_excel_publication_never_leaves_json_only_success(
    workbook_draft: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def fail(*args: object, **kwargs: object) -> None:
        raise OSError("Synthetic workbook write failure")

    monkeypatch.setattr(Path, "write_bytes", fail)
    with pytest.raises(OSError, match="workbook write"):
        check(workbook_draft)
    assert not list((workbook_draft / "outputs").glob("fact-sheet-check-*"))
    assert not list((workbook_draft / "outputs").glob(".staging-*"))


def test_control_characters_in_internal_json_never_leak_via_excel_errors(
    workbook_draft: Path,
) -> None:
    _, report = check(workbook_draft)
    fact = {**report["items"][0]["fact"], "owner": "SECRET\x01VALUE"}
    path = workbook_draft / "outputs/internal.json"
    path.write_text(
        json.dumps(
            {
                "schema_version": 1,
                "version_id": "test",
                "steward": "test",
                "facts": [fact],
            }
        )
    )
    with pytest.raises(FactSheetError, match="Excel.*character") as error:
        run_fact_sheet_check(workbook_draft, Path("outputs/internal.json"), as_of=AS_OF)
    assert "SECRET" not in str(error.value)


def test_evidence_limit_includes_unresolved_references(workbook_draft: Path) -> None:
    workbook = load_workbook(workbook_draft / DRAFT)
    evidence = workbook["Evidence"]
    row = [cell.value for cell in evidence[2]]
    row[1] = "bd_agreements/supporting_data/missing.pdf"
    for _ in range(20):
        evidence.append(row)
    workbook.save(workbook_draft / DRAFT)
    workbook.close()
    with pytest.raises(FactSheetError, match="Evidence.*20"):
        check(workbook_draft)
