import hashlib
import json
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from bd_packet_builder.cli import main
from bd_packet_builder.fact_sheet import FactSheetError
from bd_packet_builder.poc import MAINTAINED_WORKBOOK, initialize_poc, run_poc

AS_OF = date(2026, 10, 21)


@pytest.fixture(autouse=True)
def supporting_directory(tmp_path: Path) -> None:
    (tmp_path / "bd_agreements/supporting_data").mkdir(parents=True)


def fill_workbook(root: Path, value: str = "Maintained Example LLC") -> Path:
    path = root / MAINTAINED_WORKBOOK
    workbook = load_workbook(path)
    workbook["Details"]["B4"] = "Business Owner"
    values = [
        "legal-name-1",
        "firm.legal_name",
        "firm_identity",
        "text",
        value,
        None,
        "Business Owner",
        "Business Owner",
        "2026-10-01",
        "2027-10-01",
        "2026-01-01",
        None,
    ]
    for column, value in enumerate(values, 1):
        workbook["Facts"].cell(2, column, value)
    workbook.save(path)
    workbook.close()
    return path


def test_initialization_creates_one_empty_workbook_without_inventing_facts(tmp_path: Path) -> None:
    path = initialize_poc(tmp_path)
    assert path == tmp_path / MAINTAINED_WORKBOOK
    workbook = load_workbook(path)
    assert workbook["Details"]["B4"].value is None
    assert all(workbook["Facts"].cell(row, 1).value is None for row in range(2, 502))
    assert "only workbook" in workbook["Instructions"]["A2"].value
    workbook.close()
    before = path.read_bytes()
    with pytest.raises(FactSheetError, match="already exists"):
        initialize_poc(tmp_path)
    assert path.read_bytes() == before


def test_repeated_runs_read_edits_without_creating_another_maintained_workbook(
    tmp_path: Path,
) -> None:
    initialize_poc(tmp_path)
    path = fill_workbook(tmp_path)
    before = path.read_bytes()
    first = run_poc(tmp_path, as_of=AS_OF)
    assert path.read_bytes() == before
    assert (first / "input.snapshot.xlsx").read_bytes() == before
    assert not (first / "fact-sheet.xlsx").exists()
    data = json.loads((first / "fact-sheet-check.json").read_text())
    assert data["draft"]["sha256"] == hashlib.sha256(before).hexdigest()
    assert data["items"][0]["fact"]["value"]["value"] == "Maintained Example LLC"
    assert data["items"][0]["check"]["issues"] == ["missing_evidence"]
    assert "missing evidence" in (first / "validation.html").read_text().lower()
    assert "No packet" in (first / "validation.html").read_text()
    workbook = load_workbook(path)
    workbook["Facts"]["E2"] = "Edited Example LLC"
    workbook.save(path)
    workbook.close()
    second = run_poc(tmp_path, as_of=AS_OF)
    assert second != first
    latest = json.loads((second / "fact-sheet-check.json").read_text())
    assert latest["items"][0]["fact"]["value"]["value"] == "Edited Example LLC"
    assert latest["draft"]["sha256"] != data["draft"]["sha256"]
    assert list(path.parent.glob("*.xlsx")) == [path]
    assert not list((tmp_path / "outputs").rglob("synthetic*"))


def test_poc_never_searches_demo_directories_or_silently_initializes(tmp_path: Path) -> None:
    with pytest.raises(FactSheetError, match="poc-init"):
        run_poc(tmp_path, as_of=AS_OF)
    assert not (tmp_path / MAINTAINED_WORKBOOK).exists()


def test_empty_template_is_not_a_successful_application_run(tmp_path: Path) -> None:
    initialize_poc(tmp_path)
    with pytest.raises(FactSheetError):
        run_poc(tmp_path, as_of=AS_OF)
    assert not list((tmp_path / "outputs").glob("fact-sheet-check-*"))


@pytest.mark.parametrize("component", ["outputs", "outputs/maintained"])
def test_initialization_rejects_symlink_directories(tmp_path: Path, component: str) -> None:
    outside = tmp_path / "outside"
    outside.mkdir()
    target = tmp_path / component
    target.parent.mkdir(parents=True, exist_ok=True)
    target.symlink_to(outside, target_is_directory=True)
    with pytest.raises(FactSheetError, match="symlink"):
        initialize_poc(tmp_path)
    assert not list(outside.iterdir())


def test_poc_commands_announce_fixed_workbook_and_report_not_another_copy(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    assert main(["poc-init", "--root", str(tmp_path)]) == 0
    assert "only workbook" in capsys.readouterr().out
    fill_workbook(tmp_path)
    assert main(["poc", "--root", str(tmp_path), "--as-of", AS_OF.isoformat()]) == 0
    output = capsys.readouterr().out
    assert "validation.html" in output
    assert "No packet" in output
    assert "Maintained Example LLC" not in output


def test_report_escapes_untrusted_fact_keys_and_ids(tmp_path: Path) -> None:
    initialize_poc(tmp_path)
    path = fill_workbook(tmp_path)
    workbook = load_workbook(path)
    workbook["Facts"]["E2"] = "<script>alert(1)</script>"
    workbook.save(path)
    workbook.close()
    result = run_poc(tmp_path, as_of=AS_OF)
    assert "<script>" not in (result / "validation.html").read_text()


def test_maintained_workbook_resolves_native_evidence_during_intake(tmp_path: Path) -> None:
    from reportlab.pdfgen.canvas import Canvas

    from bd_packet_builder.evidence import extract_evidence
    from bd_packet_builder.inventory import build_inventory

    initialize_poc(tmp_path)
    path = fill_workbook(tmp_path)
    source = tmp_path / "bd_agreements/supporting_data/example.pdf"
    canvas = Canvas(str(source))
    canvas.drawString(40, 700, "Maintained Example LLC")
    canvas.save()
    page = extract_evidence(tmp_path, build_inventory(tmp_path)).documents[0].pages[0]
    workbook = load_workbook(path)
    workbook["Evidence"].append(
        [
            "legal-name-1",
            "bd_agreements/supporting_data/example.pdf",
            1,
            page.text,
        ]
    )
    workbook.save(path)
    workbook.close()
    result = run_poc(tmp_path, as_of=AS_OF)
    data = json.loads((result / "fact-sheet-check.json").read_text())
    assert data["facts_with_exceptions"] == 0
    assert data["items"][0]["check"]["resolved_citations"] == 1
    assert data["items"][0]["check"]["usable_for_filling"] is False


def test_failed_snapshot_write_never_publishes_a_successful_run(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    initialize_poc(tmp_path)
    path = fill_workbook(tmp_path)
    before = path.read_bytes()

    def fail(*args: object, **kwargs: object) -> None:
        raise OSError("Synthetic snapshot failure")

    monkeypatch.setattr(Path, "write_bytes", fail)
    with pytest.raises(OSError, match="snapshot failure"):
        run_poc(tmp_path, as_of=AS_OF)
    assert path.read_bytes() == before
    assert not list((tmp_path / "outputs").glob("fact-sheet-check-*"))
    assert not list((tmp_path / "outputs").glob(".staging-*"))


@pytest.mark.parametrize("value", ["20261021", "invalid", "2026-W43-3"])
def test_poc_date_errors_are_explicit(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
    value: str,
) -> None:
    assert main(["poc", "--root", str(tmp_path), "--as-of", value]) == 1
    assert "YYYY-MM-DD" in capsys.readouterr().err


def test_poc_defaults_to_today_without_a_separate_validation_command(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    initialize_poc(tmp_path)
    fill_workbook(tmp_path)
    assert main(["poc", "--root", str(tmp_path)]) == 0
    capsys.readouterr()
    report = next((tmp_path / "outputs").glob("fact-sheet-check-*/fact-sheet-check.json"))
    assert json.loads(report.read_text())["as_of"] == date.today().isoformat()
