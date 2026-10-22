"""Human-editable Excel drafts; all imported values remain unapproved proposals."""

import json
from datetime import datetime
from io import BytesIO
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import IllegalCharacterError, InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import ValidationError

from bd_packet_builder.contracts import Citation
from bd_packet_builder.evidence import extract_evidence
from bd_packet_builder.fact_sheet import FactCheck, FactSheet, FactSheetError
from bd_packet_builder.inventory import Inventory

FACT_HEADERS = (
    "Record ID",
    "Fact key",
    "Category",
    "Value type",
    "Value",
    "Signer role",
    "Owner",
    "Reviewed by",
    "Reviewed on",
    "Review due",
    "Effective from",
    "Effective to",
)
EVIDENCE_HEADERS = ("Record ID", "Supporting document", "Page", "Exact quotation", "Source SHA256")
DETAIL_FIELDS = ("schema_version", "version_id", "steward", "supersedes_sha256")
type EvidenceRow = tuple[str, str, int, str, str]


def _text(sheet: Worksheet, row: int, column: int, *, optional: bool = False) -> str:
    cell = sheet.cell(row, column)
    value = cell.value
    if value is None and optional:
        return ""
    if cell.data_type == "f":
        raise FactSheetError(f"{sheet.title}!{cell.coordinate}: formulas are not allowed")
    if not isinstance(value, str) or not value.strip() or len(value) > 20_000:
        raise FactSheetError(f"{sheet.title}!{cell.coordinate}: expected bounded text")
    return value


def _rows(sheet: Worksheet, headers: tuple[str, ...], limit: int) -> list[int]:
    if sheet.max_row > limit + 1 or sheet.max_column > len(headers):
        raise FactSheetError(f"{sheet.title}: worksheet exceeds its row/column limits")
    if tuple(sheet.cell(1, column).value for column in range(1, len(headers) + 1)) != headers:
        raise FactSheetError(f"{sheet.title}: restore the original column headings")
    return [
        row
        for row in range(2, sheet.max_row + 1)
        if any(sheet.cell(row, column).value is not None for column in range(1, len(headers) + 1))
    ]


def _date(sheet: Worksheet, row: int, column: int, *, optional: bool = False) -> str | None:
    cell = sheet.cell(row, column)
    if isinstance(cell.value, datetime):
        if cell.value.time() != datetime.min.time():
            raise FactSheetError(f"{sheet.title}!{cell.coordinate}: dates must not include a time")
        return cell.value.date().isoformat()
    value = _text(sheet, row, column, optional=optional)
    return value or None


def import_workbook(
    data: bytes,
    root: Path,
    inventory: Inventory,
) -> tuple[FactSheet, dict[str, tuple[str, ...]], list[EvidenceRow]]:
    """Resolve human-entered path/page/quote references, never guess ambiguous matches."""
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 100 or sum(entry.file_size for entry in entries) > 20 * 1024 * 1024:
                raise FactSheetError("Excel workbook exceeds its expanded size/entry limits")
            if any(
                entry.filename.endswith(".bin") or "externalLinks/" in entry.filename
                for entry in entries
            ):
                raise FactSheetError("Excel workbook must not contain macros or external links")
        workbook = load_workbook(BytesIO(data), data_only=False, keep_links=False)
    except (BadZipFile, InvalidFileException, KeyError, ValueError, ParseError, EOFError) as error:
        if isinstance(error, FactSheetError):
            raise
        raise FactSheetError("Invalid Excel workbook; use the generated .xlsx template") from None
    try:
        if not {"Instructions", "Details", "Facts", "Evidence"} <= set(workbook.sheetnames):
            raise FactSheetError(
                "Excel workbook requires Instructions, Details, Facts and Evidence"
            )
        if set(workbook.sheetnames) - {"Instructions", "Details", "Facts", "Evidence", "Checks"}:
            raise FactSheetError("Excel workbook contains unexpected worksheets")
        details = workbook["Details"]
        rows = _rows(details, ("Field", "Value"), 4)
        if (
            rows != [2, 3, 4, 5]
            or tuple(details.cell(row, 1).value for row in rows) != DETAIL_FIELDS
        ):
            raise FactSheetError("Details: restore the original metadata fields")
        if _text(details, 2, 2) != "1":
            raise FactSheetError("Details!B2: unsupported workbook schema version")
        metadata = {
            "schema_version": 1,
            "version_id": _text(details, 3, 2),
            "steward": _text(details, 4, 2),
            "supersedes_sha256": _text(details, 5, 2, optional=True) or None,
        }
        facts = workbook["Facts"]
        records: list[dict[str, object]] = []
        for row in _rows(facts, FACT_HEADERS, 500):
            records.append(
                {
                    "fact_id": _text(facts, row, 1),
                    "key": _text(facts, row, 2),
                    "category": _text(facts, row, 3),
                    "value": {"kind": _text(facts, row, 4), "value": _text(facts, row, 5)},
                    "role": _text(facts, row, 6, optional=True) or None,
                    "owner": _text(facts, row, 7),
                    "reviewed_by": _text(facts, row, 8),
                    "reviewed_on": _date(facts, row, 9),
                    "review_due": _date(facts, row, 10),
                    "effective_from": _date(facts, row, 11),
                    "effective_to": _date(facts, row, 12, optional=True),
                    "citations": [],
                }
            )
        try:
            sheet = FactSheet.model_validate_json(json.dumps({**metadata, "facts": records}))
        except ValidationError:
            raise FactSheetError(
                "Facts/Details: invalid identifiers, types, roles or ISO review/effective dates"
            ) from None
        evidence_sheet = workbook["Evidence"]
        evidence_rows: list[EvidenceRow] = []
        reference_counts = dict.fromkeys((fact.fact_id for fact in sheet.facts), 0)
        for row in _rows(evidence_sheet, EVIDENCE_HEADERS, 10_000):
            fact_id = _text(evidence_sheet, row, 1)
            if fact_id not in reference_counts:
                raise FactSheetError(f"Evidence!A{row}: no matching Facts record ID")
            reference_counts[fact_id] += 1
            if reference_counts[fact_id] > 20:
                raise FactSheetError(
                    f"Evidence!A{row}: at most 20 references per fact are supported"
                )
            page = evidence_sheet.cell(row, 3).value
            if type(page) is not int or page < 1:
                raise FactSheetError(f"Evidence!C{row}: expected a positive integer page")
            evidence_rows.append(
                (
                    fact_id,
                    _text(evidence_sheet, row, 2),
                    page,
                    _text(evidence_sheet, row, 4),
                    _text(evidence_sheet, row, 5, optional=True),
                )
            )
        bundle = extract_evidence(root, inventory) if evidence_rows else None
        by_path = {source.path: source for source in inventory.sources}
        citations: dict[str, list[Citation]] = {fact.fact_id: [] for fact in sheet.facts}
        issues: dict[str, list[str]] = {fact.fact_id: [] for fact in sheet.facts}
        resolved_rows: list[EvidenceRow] = []
        for fact_id, path, page_number, quote, digest in evidence_rows:
            source = by_path.get(path)
            page_text = (
                next(
                    (
                        page.text
                        for doc in bundle.documents
                        for page in doc.pages
                        if source is not None
                        and page.source_id == source.source_id
                        and page.page == page_number
                    ),
                    "",
                )
                if bundle is not None
                else ""
            )
            if source is None or not page_text:
                issues[fact_id].append("citation_unresolved")
            elif digest and digest != source.sha256:
                issues[fact_id].append("citation_source_version_mismatch")
            elif page_text.count(quote) != 1:
                issues[fact_id].append("citation_quote_missing_or_ambiguous")
            else:
                start = page_text.index(quote)
                citations[fact_id].append(
                    Citation(
                        source_id=source.source_id,
                        page=page_number,
                        start=start,
                        end=start + len(quote),
                        quote=quote,
                    )
                )
                digest = source.sha256
            resolved_rows.append((fact_id, path, page_number, quote, digest))
        records = [
            {
                **fact.model_dump(mode="json"),
                "citations": [
                    citation.model_dump(mode="json") for citation in citations[fact.fact_id]
                ],
            }
            for fact in sheet.facts
        ]
        try:
            sheet = FactSheet.model_validate_json(json.dumps({**metadata, "facts": records}))
        except ValidationError:
            raise FactSheetError(
                "Evidence: citations exceed the supported count or quote limits"
            ) from None
        return (
            sheet,
            {key: tuple(dict.fromkeys(value)) for key, value in issues.items()},
            resolved_rows,
        )
    finally:
        workbook.close()


def _write_table(
    workbook: Workbook,
    title: str,
    rows: list[tuple[str | int | None, ...]],
) -> Worksheet:
    sheet: Worksheet = workbook.create_sheet(title)
    for row in rows:
        try:
            sheet.append(row)
        except IllegalCharacterError:
            raise FactSheetError(
                f"{title}: Excel cannot represent a supplied control character"
            ) from None
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
                cell.number_format = "@"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
    for column in range(1, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 24
    sheet.freeze_panes = "A2"
    return sheet


def export_workbook(
    sheet: FactSheet | None,
    inventory: Inventory,
    *,
    checks: tuple[FactCheck, ...] = (),
    evidence_rows: list[EvidenceRow] | None = None,
    as_of: str | None = None,
) -> bytes:
    """Create a formatted editable workbook; even formula-like text is literal text."""
    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    instructions = _write_table(
        workbook,
        "Instructions",
        [
            ("DRAFT FACT SHEET - NOT APPROVED",),
            (
                "This is the only workbook you maintain: edit Facts, Details and Evidence, "
                "check your changes, and save as .xlsx. No code or terminal is required.",
            ),
            (
                "The business owner is responsible for factual accuracy. "
                "No separate auditor or second-person fact-sheet approval is required.",
            ),
            (
                "Reviewed by means the person who checked the entry; "
                "this can be the same person as Owner.",
            ),
            (
                "The application cannot reliably detect a plausible typo. "
                "Check names, identifiers and values before saving.",
            ),
            (
                "Save this workbook before the PoC is run. The application validates it "
                "during intake; archived snapshots are not workbooks to maintain.",
            ),
            (
                "First use: enter your name in Details > steward and add at least one "
                "Facts row. The blank template contains no firm facts.",
            ),
            (
                "For normal PoC runs, read validation.html in the run results. "
                "The application does not write check results back into this maintained workbook.",
            ),
            (
                "Never apply signatures/dates or infer missing facts. "
                "No row is approved for filling.",
            ),
            ("Facts: keep Record ID unique; use the same Fact key for versions of one fact.",),
            ("Value must be TEXT (including CRDs). Use text or identifier; never formulas.",),
            (
                "Dates: YYYY-MM-DD text or Excel date cells, without times. "
                "End/due dates are inclusive.",
            ),
            (
                "Evidence: one row per reference. Enter Record ID, "
                "supporting document path, page, exact quote.",
            ),
            (
                "Paths are relative to your data root under bd_agreements/supporting_data/; "
                "not completed forms.",
            ),
            (
                "Do not edit hidden technical columns. "
                "Ask the technical operator for help if a supporting document changes.",
            ),
            (
                "Owner/reviewer names record responsibility, not an independent audit. "
                "Review on change and at least annually.",
            ),
            (
                "Checks is a point-in-time report ONLY. "
                "Editing this workbook makes those checks stale.",
            ),
            (
                f"Last evaluated: {as_of or 'Not yet evaluated'}; "
                "semantic support and approval UNVERIFIED.",
            ),
            (
                "Workbook, quotes and internal reports may be confidential. "
                "Do not send with issuer packets.",
            ),
            (
                "Category: firm_identity with firm.* key; "
                "signer_role with signer_roles.* key and named person.",
            ),
        ],
    )
    instructions.column_dimensions["A"].width = 110
    _write_table(
        workbook,
        "Details",
        [
            ("Field", "Value"),
            ("schema_version", "1"),
            ("version_id", sheet.version_id if sheet else "draft-001"),
            ("steward", sheet.steward if sheet else None),
            ("supersedes_sha256", sheet.supersedes_sha256 if sheet else None),
        ],
    )
    facts = _write_table(
        workbook,
        "Facts",
        [
            FACT_HEADERS,
            *(
                (
                    fact.fact_id,
                    fact.key,
                    fact.category,
                    fact.value.kind,
                    fact.value.value,
                    fact.role,
                    fact.owner,
                    fact.reviewed_by,
                    fact.reviewed_on.isoformat(),
                    fact.review_due.isoformat(),
                    fact.effective_from.isoformat(),
                    fact.effective_to.isoformat() if fact.effective_to else None,
                )
                for fact in (sheet.facts if sheet else ())
            ),
        ],
    )
    facts.freeze_panes = "F2"
    facts.column_dimensions["E"].width = 44
    facts.auto_filter.ref = facts.dimensions
    for column, choices in [("C", '"firm_identity,signer_role"'), ("D", '"text,identifier"')]:
        validation = DataValidation(type="list", formula1=choices, allow_blank=False)
        validation.errorTitle = "Choose a supported value"
        validation.error = "Use a value from the dropdown."
        validation.showErrorMessage = True
        facts.add_data_validation(validation)
        validation.add(f"{column}2:{column}501")
    # New identifier cells must stay text when edited in Excel, not lose leading zeros.
    for row in range(2, 502):
        facts.cell(row, 5).number_format = "@"
    if evidence_rows is None:
        sources = {source.source_id: source for source in inventory.sources}
        evidence_rows = []
        for fact in sheet.facts if sheet else ():
            for citation in fact.citations:
                source = sources.get(citation.source_id)
                evidence_rows.append(
                    (
                        fact.fact_id,
                        source.path if source else f"Unresolved source ID: {citation.source_id}",
                        citation.page,
                        citation.quote,
                        source.sha256 if source else "",
                    )
                )
    evidence = _write_table(workbook, "Evidence", [EVIDENCE_HEADERS, *evidence_rows])
    evidence.column_dimensions["B"].width = 55
    evidence.column_dimensions["D"].width = 80
    evidence.column_dimensions["E"].hidden = True
    evidence.auto_filter.ref = evidence.dimensions
    _write_table(
        workbook,
        "Checks",
        [
            (
                "Record ID",
                "Mechanical checks (NOT APPROVAL)",
                "Resolved citations",
                "Evaluated as of",
            ),
            *(
                (
                    check.fact_id,
                    "; ".join(check.issues) or "No mechanical exceptions; NOT APPROVED",
                    check.resolved_citations,
                    as_of,
                )
                for check in checks
            ),
        ],
    )
    result = BytesIO()
    try:
        workbook.save(result)
        return result.getvalue()
    finally:
        workbook.close()
